import re
import warnings
import pandas as pd
from netmiko import ConnectHandler
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings(action='ignore', module='.*paramiko.*')

def extraer_ipam(dispositivo):
    print(f"Intentando entrar a {dispositivo['name']} ({dispositivo['ip']})...")
    try:
        params = dispositivo.copy()
        nombre_router = params.pop('name')
        
        conexion = ConnectHandler(**params)
        salida = conexion.send_command("show ip interface brief")
        conexion.disconnect()
        
        patron = r"([A-Za-z0-9/]+)\s+([0-9]+\.[0-9]+\.[0-9]+\.[0-9]+)\s+\w+\s+\w+\s+(up|down|administratively down)\s+(up|down)"
        coincidencias = re.findall(patron, salida)
        
        lista_interfaces = []
        for interfaz, ip, estado_fisico, protocolo in coincidencias:
            tipo_interfaz = "Virtual (Loopback)" if "Loopback" in interfaz else "Física"
            lista_interfaces.append({
                "Interfaz": interfaz,
                "Dirección IP": ip,
                "Tipo": tipo_interfaz,
                "Estado Físico": estado_fisico,
                "Estado Protocolo": protocolo
            })
        print(f"¡Listo! Se extrajeron las IPs de {nombre_router}.")
        return nombre_router, lista_interfaces
    except Exception as e:
        print(f"Uy, falló la conexión con {dispositivo['name']}. Saltando al siguiente...")
        return dispositivo['name'], []

def dar_estilo_excel(archivo):
    import openpyxl
    wb = openpyxl.load_workbook(archivo)
    
    color_encabezado = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fuente_datos = Font(name="Calibri", size=11)
    centrar = Alignment(horizontal="center", vertical="center")
    alinear_izq = Alignment(horizontal="left", vertical="center")
    borde_suave = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    for hoja in wb.worksheets:
        # Congelar el encabezado para que siempre se vea al bajar
        hoja.freeze_panes = "A2"
        
        for col in hoja.iter_cols(min_row=1, max_row=hoja.max_row, min_col=1, max_col=hoja.max_column):
            ancho_max = max(len(str(celda.value or '')) for celda in col)
            letra_col = get_column_letter(col[0].column)
            hoja.column_dimensions[letra_col].width = max(ancho_max + 4, 15)
            
        for celda in hoja[1]:
            celda.fill = color_encabezado
            celda.font = fuente_encabezado
            celda.alignment = centrar
            
        for fila in hoja.iter_rows(min_row=2, max_row=hoja.max_row):
            for celda in fila:
                celda.font = fuente_datos
                celda.border = borde_suave
                celda.alignment = centrar if celda.column != 1 else alinear_izq
                
    wb.save(archivo)

def main():
    print("Empezando a sacar las IPs de los routers del equipo 5...")
    
    routers = [
        {'device_type': 'cisco_ios', 'ip': '192.168.1.22', 'username': 'admin', 'password': 'cisco', 'name': 'R1-5'},
        {'device_type': 'cisco_ios', 'ip': '10.0.5.2',   'username': 'admin', 'password': 'cisco', 'name': 'R2-5'},
        {'device_type': 'cisco_ios', 'ip': '10.0.5.6',   'username': 'admin', 'password': 'cisco', 'name': 'R3-5'}
    ]
    
    resultados = {}
    for router in routers:
        nombre, datos = extraer_ipam(router)
        if datos:
            resultados[nombre] = datos
            
    if resultados:
        nombre_archivo = "Reporte_IPAM_Equipo_5.xlsx"
        with pd.ExcelWriter(nombre_archivo, engine='openpyxl') as writer:
            for nombre_router, info_router in resultados.items():
                df = pd.DataFrame(info_router)
                df.to_excel(writer, sheet_name=nombre_router, index=False)
                
        dar_estilo_excel(nombre_archivo)
        print(f"\n¡Terminamos! El reporte quedó súper bien y se guardó como {nombre_archivo}")
    else:
        print("\nNo se pudo sacar información de ningún equipo.")

if __name__ == "__main__":
    main()
